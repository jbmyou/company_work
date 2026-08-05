from pathlib import Path
import shutil

from openpyxl import load_workbook


src_dir = r"D:\1.다운로드\신탁통지서"
key_list = "./260723_P15_신탁.xlsx"  # 채무자키, 계좌키 중 채무자키


# src_dir 폴더에 있는 파일 중 파일명이 key_list에 있는 채무자키로 시작하는 파일은 "신탁" 폴더를 만들어서 그 안으로 옮기고,
# key_list에 각각 몇 개의 파일이 옮겨졌는지 열을 추가하여 기록


DEST_DIR_NAME = "신탁"
KEY_COLUMN_NAME = "채무자키"
COUNT_COLUMN_NAME = "옮긴파일수"


def normalize_key(value):
    """엑셀 셀 값을 파일명 비교용 문자열 키로 변환한다."""
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def unique_destination_path(dest_dir, file_name):
    """같은 이름의 파일이 이미 있으면 파일명 뒤에 번호를 붙인다."""
    dest_path = dest_dir / file_name
    if not dest_path.exists():
        return dest_path

    stem = dest_path.stem
    suffix = dest_path.suffix
    index = 1

    while True:
        candidate = dest_dir / f"{stem} ({index}){suffix}"
        if not candidate.exists():
            return candidate
        index += 1


def find_column(ws, column_name, default_col=1):
    for cell in ws[1]:
        if str(cell.value).strip() == column_name:
            return cell.column
    return default_col


def find_or_create_column(ws, column_name):
    for cell in ws[1]:
        if str(cell.value).strip() == column_name:
            return cell.column

    col = ws.max_column + 1
    ws.cell(row=1, column=col, value=column_name)
    return col


def main():
    source_dir = Path(src_dir)
    excel_path = Path(key_list)
    dest_dir = source_dir / DEST_DIR_NAME

    if not source_dir.exists():
        raise FileNotFoundError(f"소스 폴더를 찾을 수 없습니다: {source_dir}")

    if not excel_path.exists():
        raise FileNotFoundError(f"엑셀 파일을 찾을 수 없습니다: {excel_path}")

    wb = load_workbook(excel_path)
    ws = wb.active

    key_col = find_column(ws, KEY_COLUMN_NAME, default_col=1)
    count_col = find_or_create_column(ws, COUNT_COLUMN_NAME)

    row_keys = {}
    keys = set()

    for row in range(2, ws.max_row + 1):
        key = normalize_key(ws.cell(row=row, column=key_col).value)
        row_keys[row] = key
        if key:
            keys.add(key)

    counts = {key: 0 for key in keys}
    sorted_keys = sorted(keys, key=len, reverse=True)

    dest_dir.mkdir(parents=True, exist_ok=True)

    for file_path in source_dir.iterdir():
        if not file_path.is_file():
            continue

        matched_key = next(
            (key for key in sorted_keys if file_path.name.startswith(key)),
            None,
        )

        if matched_key is None:
            continue

        dest_path = unique_destination_path(dest_dir, file_path.name)
        shutil.move(str(file_path), str(dest_path))
        counts[matched_key] += 1

    for row, key in row_keys.items():
        ws.cell(row=row, column=count_col, value=counts.get(key, 0))

    wb.save(excel_path)

    total_count = sum(counts.values())
    print(f"완료: {total_count}개 파일을 '{dest_dir}' 폴더로 옮겼습니다.")
    print(f"엑셀에 '{COUNT_COLUMN_NAME}' 열을 기록했습니다: {excel_path}")


if __name__ == "__main__":
    main()
